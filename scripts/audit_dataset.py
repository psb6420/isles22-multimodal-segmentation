#!/usr/bin/env python3
"""Audit the extracted ISLES'22 training dataset without modifying it."""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

import nibabel as nib
import numpy as np
from openpyxl import load_workbook
from scipy import ndimage


MODALITIES = ("dwi", "adc", "flair", "mask")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--center-file", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def load_center_map(path: Path) -> dict[str, int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    if tuple(header[:2]) != ("case", "center"):
        raise ValueError(f"Unexpected center file header: {header}")
    result = {}
    for case, center, *_ in rows:
        if case is not None and center is not None:
            result[str(case)] = int(center)
    return result


def expected_paths(root: Path, subject: str) -> dict[str, Path]:
    session = root / subject / "ses-0001"
    prefix = f"{subject}_ses-0001"
    return {
        "dwi": session / "dwi" / f"{prefix}_dwi.nii.gz",
        "adc": session / "dwi" / f"{prefix}_adc.nii.gz",
        "flair": session / "anat" / f"{prefix}_FLAIR.nii.gz",
        "mask": root / "derivatives" / subject / "ses-0001" / f"{prefix}_msk.nii.gz",
    }


def shape_text(shape: tuple[int, ...]) -> str:
    return "x".join(str(int(value)) for value in shape)


def float_tuple(values: tuple[float, ...], digits: int = 6) -> tuple[float, ...]:
    return tuple(round(float(value), digits) for value in values)


def grid_matches(first: nib.spatialimages.SpatialImage, second: nib.spatialimages.SpatialImage) -> bool:
    return first.shape == second.shape and np.allclose(first.affine, second.affine, atol=1e-4, rtol=0)


def image_metadata(image: nib.spatialimages.SpatialImage) -> dict[str, object]:
    header = image.header
    return {
        "shape": shape_text(image.shape),
        "spacing": float_tuple(header.get_zooms()[:3]),
        "orientation": "".join(nib.aff2axcodes(image.affine)),
        "dtype": str(header.get_data_dtype()),
        "qform_code": int(header["qform_code"]),
        "sform_code": int(header["sform_code"]),
    }


def intensity_stats(image: nib.spatialimages.SpatialImage) -> dict[str, object]:
    data = np.asanyarray(image.dataobj)
    finite = np.isfinite(data)
    finite_values = data[finite]
    if finite_values.size == 0:
        return {"nan_voxels": int(np.isnan(data).sum()), "inf_voxels": int(np.isinf(data).sum())}
    return {
        "nan_voxels": int(np.isnan(data).sum()),
        "inf_voxels": int(np.isinf(data).sum()),
        "min": float(finite_values.min()),
        "max": float(finite_values.max()),
        "nonzero_voxels": int(np.count_nonzero(finite_values)),
    }


def mask_stats(image: nib.spatialimages.SpatialImage) -> dict[str, object]:
    data = np.asanyarray(image.dataobj)
    unique = np.unique(data)
    foreground = data > 0
    structure = ndimage.generate_binary_structure(rank=3, connectivity=3)
    labeled, lesion_count = ndimage.label(foreground, structure=structure)
    component_sizes = np.bincount(labeled.ravel())[1:]
    voxel_volume_mm3 = float(np.prod(image.header.get_zooms()[:3]))
    lesion_voxels = int(foreground.sum())
    return {
        "labels": ",".join(str(float(value)).rstrip("0").rstrip(".") for value in unique),
        "lesion_voxels": lesion_voxels,
        "lesion_volume_ml": lesion_voxels * voxel_volume_mm3 / 1000.0,
        "lesion_count_26c": int(lesion_count),
        "smallest_component_voxels": int(component_sizes.min()) if component_sizes.size else 0,
        "largest_component_voxels": int(component_sizes.max()) if component_sizes.size else 0,
    }


def counter_dict(counter: Counter) -> dict[str, int]:
    return {str(key): int(value) for key, value in sorted(counter.items(), key=lambda item: str(item[0]))}


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    centers = load_center_map(args.center_file)
    subjects = sorted(path.name for path in args.data_root.glob("sub-strokecase*") if path.is_dir())

    rows: list[dict[str, object]] = []
    missing: list[dict[str, str]] = []
    counters = {
        modality: {
            "shape": Counter(),
            "spacing": Counter(),
            "orientation": Counter(),
            "dtype": Counter(),
        }
        for modality in MODALITIES
    }

    for index, subject in enumerate(subjects, start=1):
        paths = expected_paths(args.data_root, subject)
        absent = [name for name, path in paths.items() if not path.is_file()]
        if absent:
            for name in absent:
                missing.append({"subject": subject, "modality": name, "expected_path": str(paths[name])})
            continue

        images = {name: nib.load(path) for name, path in paths.items()}
        row: dict[str, object] = {"subject": subject, "center": centers.get(subject, "")}

        for name, image in images.items():
            metadata = image_metadata(image)
            for field, value in metadata.items():
                row[f"{name}_{field}"] = value
            for field in ("shape", "spacing", "orientation", "dtype"):
                counters[name][field][str(metadata[field])] += 1

        row["dwi_adc_same_grid"] = grid_matches(images["dwi"], images["adc"])
        row["dwi_mask_same_grid"] = grid_matches(images["dwi"], images["mask"])
        row["dwi_flair_same_grid"] = grid_matches(images["dwi"], images["flair"])

        for name in ("dwi", "adc", "flair"):
            for field, value in intensity_stats(images[name]).items():
                row[f"{name}_{field}"] = value
        for field, value in mask_stats(images["mask"]).items():
            row[f"mask_{field}"] = value

        rows.append(row)
        if index % 25 == 0 or index == len(subjects):
            print(f"Audited {index}/{len(subjects)} subjects", flush=True)

    fieldnames = sorted({key for row in rows for key in row}, key=lambda key: (key != "subject", key != "center", key))
    csv_path = args.output_dir / "isles22_case_audit.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    summary = {
        "data_root": str(args.data_root.resolve()),
        "subject_directories": len(subjects),
        "complete_subjects": len(rows),
        "center_ids": counter_dict(Counter(row["center"] for row in rows)),
        "missing_files": missing,
        "grid_matches": {
            "dwi_adc": int(sum(bool(row["dwi_adc_same_grid"]) for row in rows)),
            "dwi_mask": int(sum(bool(row["dwi_mask_same_grid"]) for row in rows)),
            "dwi_flair": int(sum(bool(row["dwi_flair_same_grid"]) for row in rows)),
        },
        "header_distributions": {
            modality: {field: counter_dict(counter) for field, counter in fields.items()}
            for modality, fields in counters.items()
        },
        "mask_label_sets": counter_dict(Counter(str(row["mask_labels"]) for row in rows)),
        "nonfinite_voxels": {
            modality: {
                "cases_with_nan": int(sum(int(row[f"{modality}_nan_voxels"]) > 0 for row in rows)),
                "cases_with_inf": int(sum(int(row[f"{modality}_inf_voxels"]) > 0 for row in rows)),
            }
            for modality in ("dwi", "adc", "flair")
        },
        "lesion_volume_ml": {
            "min": float(min(float(row["mask_lesion_volume_ml"]) for row in rows)),
            "median": float(np.median([float(row["mask_lesion_volume_ml"]) for row in rows])),
            "max": float(max(float(row["mask_lesion_volume_ml"]) for row in rows)),
        },
        "lesion_count_26c": {
            "min": int(min(int(row["mask_lesion_count_26c"]) for row in rows)),
            "median": float(np.median([int(row["mask_lesion_count_26c"]) for row in rows])),
            "max": int(max(int(row["mask_lesion_count_26c"]) for row in rows)),
        },
    }
    json_path = args.output_dir / "isles22_summary.json"
    json_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    print(f"Wrote {csv_path}")
    print(f"Wrote {json_path}")


if __name__ == "__main__":
    main()
